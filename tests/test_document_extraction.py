from docx import Document

from ai_engine.pdf_extractor import extract_text_from_docx, process_document
from backend.ai_client import ai_pipeline


def _make_docx(path):
    document = Document()
    document.add_paragraph("Either party may terminate this agreement with 30 days notice.")
    document.add_paragraph("Client shall pay all invoices within 30 days.")
    document.save(path)


def test_extract_text_from_docx_reads_paragraphs(tmp_path):
    docx_path = tmp_path / "contract.docx"
    _make_docx(docx_path)

    text = extract_text_from_docx(docx_path)

    assert "Either party may terminate" in text
    assert "Client shall pay" in text


def test_process_document_supports_docx(tmp_path):
    docx_path = tmp_path / "contract.docx"
    _make_docx(docx_path)

    result = process_document(docx_path)

    assert result["source_type"] == "docx"
    assert result["num_pages"] == 0
    assert "Client shall pay" in result["clean_text"]
    assert result["paragraphs"]


def test_ai_pipeline_extract_text_supports_docx(tmp_path):
    docx_path = tmp_path / "contract.docx"
    _make_docx(docx_path)

    text = ai_pipeline.extract_text(docx_path)

    assert "Either party may terminate" in text
    assert "Client shall pay" in text
